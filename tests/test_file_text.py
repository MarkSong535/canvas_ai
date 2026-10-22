"""Text extraction from course files (src/utils/file_text.py). Run with `pytest tests` or directly."""

import io
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from src.utils.file_text import extract_text  # noqa: E402


def make_pdf(text: str) -> bytes:
    """A minimal one-page PDF with a real text layer."""
    content = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>stream\n" % len(content) + content + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out, offsets = io.BytesIO(b"%PDF-1.4\n"), []
    out.seek(0, 2)
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def make_docx(paragraphs, table=None) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_text_by_page():
    text, note = extract_text(make_pdf("Question 1: Rewrite in active voice."), "worksheet.pdf")
    assert "--- Page 1 ---" in text and "Question 1: Rewrite in active voice." in text and note == ""


def test_scanned_pdf_says_why():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)
    text, note = extract_text(buffer.getvalue(), "scan.pdf")
    assert text == "" and "no text layer" in note


def test_docx_paragraphs_and_tables():
    data = make_docx(["Master the Active Voice", "1. The ball was thrown by Sam."], [["Passive", "Active"], ["was eaten", "ate"]])
    text, _ = extract_text(data, "exercise.docx")
    assert "1. The ball was thrown by Sam." in text and "Passive | Active" in text


def test_pptx_slides_and_notes():
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "A* search"
    slide.notes_slide.notes_text_frame.text = "Mention admissibility"
    buffer = io.BytesIO()
    presentation.save(buffer)
    text, _ = extract_text(buffer.getvalue(), "lecture.pptx")
    assert "--- Slide 1 ---" in text and "A* search" in text and "Notes: Mention admissibility" in text


def test_xlsx_sheets():
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Grades"
    workbook.active.append(["Lab", "Score"])
    workbook.active.append(["Lab 1", 93.5])
    buffer = io.BytesIO()
    workbook.save(buffer)
    text, _ = extract_text(buffer.getvalue(), "grades.xlsx")
    assert "--- Sheet: Grades ---" in text and "Lab 1\t93.5" in text


def test_text_html_and_notebooks():
    assert extract_text(b"def f():\n    return 1\n", "solution.py")[0].startswith("def f()")
    text, _ = extract_text(b"<html><style>x{}</style><p>Hello</p><p>World</p></html>", "page.html")
    assert "Hello" in text and "World" in text and "x{}" not in text
    notebook = json.dumps({"cells": [{"cell_type": "markdown", "source": ["# Lab 2"]}]}).encode()
    assert "# Lab 2" in extract_text(notebook, "lab2.ipynb")[0]


def test_unreadable_types_say_why():
    assert extract_text(b"\x89PNG...", "diagram.png", "image/png") == ("", extract_text(b"", "x.png", "image/png")[1])
    assert "no readable text" in extract_text(b"\x89PNG...", "diagram.png", "image/png")[1]
    text, note = extract_text(b"\xd0\xcf", "old.doc")  # not a real .doc
    assert text == "" and ".doc" in note and ("isn't installed" in note or "couldn't be read" in note)
    assert "couldn't be read" in extract_text(b"not a pdf", "broken.pdf")[1]


if __name__ == "__main__":
    for name, fn in [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_")]:
        fn()
        print(f"PASS {name}")
